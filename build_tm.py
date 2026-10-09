#!/usr/bin/env python3
"""
Build tm-<locale>.json glossary files for the Intenta Funnels Localization plugin.

Usage:
    python3 build_tm.py <glossary file> <locale> [output folder]

    <glossary file>  .csv, .xlsx or .numbers
    <locale>         es | ptbr
    [output folder]  where tm-<locale>.json is written (default: next to this script)

Examples:
    python3 build_tm.py Intenta_ES_glossary.xlsx es
    python3 build_tm.py Intenta_PTBR_glossary.csv ptbr ~/intenta-funnels-tm

Expected format: first row is a header, one row per term, with
    • an English column  — header contains "english", "source" or "en"
    • a target column    — header contains "spanish"/"español"/"es" or
                           "portuguese"/"português"/"pt"/"target"
If the headers aren't recognized, the first two columns are used (EN, target).

What it does:
    1. Reads every EN → target pair, trimming whitespace and skipping empty rows
    2. If the same EN has several translations, keeps the most frequent one
       (and lists the conflicts so you can fix them at the source)
    3. Writes a flat JSON object {"English": "Translation", ...}, one entry per
       line in original order, ready to upload to the intenta-funnels-tm repo

Requirements: .xlsx needs `pip3 install openpyxl`; .numbers needs `pip3 install numbers-parser`.
"""

import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

VALID_LOCALES = {"es", "ptbr"}

EN_HINTS = ["english", "source", "en"]
TARGET_HINTS = {
    "es": ["spanish", "español", "espanol", "es", "target", "translation"],
    "ptbr": ["portuguese", "português", "portugues", "pt-br", "ptbr", "pt", "target", "translation"],
}


def read_rows(path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = path.read_bytes().decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(text.splitlines(), dialect)]
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("ERROR: openpyxl not installed. Run: pip3 install openpyxl")
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if suffix == ".numbers":
        try:
            from numbers_parser import Document
        except ImportError:
            sys.exit("ERROR: numbers-parser not installed. Run: pip3 install numbers-parser")
        table = Document(str(path)).sheets[0].tables[0]
        return [list(r) for r in table.rows(values_only=True)]
    sys.exit(f"ERROR: unsupported file type {suffix!r} (use .csv, .xlsx or .numbers)")


def find_column(header, hints, exclude=None):
    # Exact header match first, then substring match (hints longer than 2 chars only)
    for exact in (True, False):
        for i, h in enumerate(header):
            if i == exclude:
                continue
            for hint in hints:
                if (h == hint) if exact else (len(hint) > 2 and hint in h):
                    return i
    return None


def main():
    if len(sys.argv) not in (3, 4):
        print(__doc__)
        sys.exit(1)

    src = Path(sys.argv[1]).expanduser()
    locale = sys.argv[2].lower().replace("-", "")
    out_dir = Path(sys.argv[3]).expanduser() if len(sys.argv) == 4 else Path(__file__).resolve().parent

    if not src.exists():
        sys.exit(f"ERROR: file not found — {src}")
    if locale not in VALID_LOCALES:
        sys.exit(f"ERROR: unknown locale {locale!r} (expected one of {sorted(VALID_LOCALES)})")

    rows = read_rows(src)
    if len(rows) < 2:
        sys.exit("ERROR: file has no data rows")

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    en_col = find_column(header, EN_HINTS)
    if en_col is None:
        en_col = 0
    tgt_col = find_column(header, TARGET_HINTS[locale], exclude=en_col)
    if tgt_col is None:
        tgt_col = 1 if en_col == 0 else 0
    print(f"Header: {rows[0]}")
    print(f"Using EN column {en_col} ({rows[0][en_col]!r}), target column {tgt_col} ({rows[0][tgt_col]!r})")

    en_to_targets = defaultdict(Counter)
    order = []
    skipped = 0
    for r in rows[1:]:
        if len(r) <= max(en_col, tgt_col):
            skipped += 1
            continue
        en, tgt = r[en_col], r[tgt_col]
        en = str(en).strip() if en is not None else ""
        tgt = str(tgt).strip() if tgt is not None else ""
        if not en or not tgt:
            skipped += 1
            continue
        if en not in en_to_targets:
            order.append(en)
        en_to_targets[en][tgt] += 1

    tm, conflicts = {}, []
    for en in order:
        targets = en_to_targets[en]
        if len(targets) > 1:
            conflicts.append((en, dict(targets)))
        tm[en] = targets.most_common(1)[0][0]

    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"tm-{locale}.json"
    with out_path.open("w", encoding="utf-8") as f:
        f.write("{\n")
        items = list(tm.items())
        for i, (k, v) in enumerate(items):
            comma = "," if i < len(items) - 1 else ""
            f.write(f"{json.dumps(k, ensure_ascii=False)}: {json.dumps(v, ensure_ascii=False)}{comma}\n")
        f.write("}\n")

    print(f"\n✓ Wrote {len(tm)} entries to {out_path}")
    print(f"  Skipped empty/malformed rows: {skipped}")
    print(f"  English terms with several translations (kept most common): {len(conflicts)}")
    for en, targets in conflicts[:10]:
        print(f"  {en!r}")
        for tgt, count in sorted(targets.items(), key=lambda x: -x[1]):
            print(f"      [{count}×] {tgt!r}{' ← kept' if tgt == tm[en] else ''}")


if __name__ == "__main__":
    main()
